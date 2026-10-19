from __future__ import annotations

import argparse
import hashlib
import json
import re
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_CASE_XLSX = Path(r"D:\qq文件\人工智能OPC案例汇总_0923新增.xlsx")
DEFAULT_UNIVERSITY_XLSX = Path(r"D:\qq文件\高校OPC数据集_0923.xlsx")
DEFAULT_OUT = Path("outputs/data-import-20260929")
BATCH = "data-import-20260929"
CASE_REVIEWER = "excel-import-20260929"
ACCESS_DATE = "2026-09-29"
PROVINCE_ALIASES = {
    "北京": "北京市", "天津": "天津市", "上海": "上海市", "重庆": "重庆市",
    "河北": "河北省", "山西": "山西省", "辽宁": "辽宁省", "吉林": "吉林省",
    "黑龙江": "黑龙江省", "江苏": "江苏省", "浙江": "浙江省", "安徽": "安徽省",
    "福建": "福建省", "江西": "江西省", "山东": "山东省", "河南": "河南省",
    "湖北": "湖北省", "湖南": "湖南省", "广东": "广东省", "海南": "海南省",
    "四川": "四川省", "贵州": "贵州省", "云南": "云南省", "陕西": "陕西省",
    "甘肃": "甘肃省", "青海": "青海省", "台湾": "台湾省", "内蒙古": "内蒙古自治区",
    "广西": "广西壮族自治区", "西藏": "西藏自治区", "宁夏": "宁夏回族自治区",
    "新疆": "新疆维吾尔自治区",
}


@dataclass
class CaseRow:
    excel_row: int
    article_title: str
    title: str
    source_url: str
    province: str | None
    city: str | None
    district: str | None
    actor_name: str | None
    student_startup: str | None
    category: str
    subcategory: str
    publish_date: str | None
    summary: str
    ai_tools: str | None
    outcome: str | None


@dataclass
class UniversityRow:
    record_type: str
    record_type_label: str
    record_code: str
    record_name: str
    institution_name: str | None
    province: str | None
    city: str | None
    district: str | None
    verification_status: str
    evidence_grade: str | None
    date_original: str | None
    source_title: str | None
    source_url: str | None
    summary: str | None
    notes: str | None
    raw: dict[str, Any]


def clean(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = re.sub(r"\s+", " ", str(value).strip())
    if text in {"—", "-", "无", "未明确", "不详"}:
        return None
    return text or None


def normalize_province(value: str | None) -> str | None:
    if not value:
        return None
    return PROVINCE_ALIASES.get(value, value)


def parse_date(value: Any) -> str | None:
    text = clean(value)
    if not text:
        return None
    match = re.search(r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})", text)
    if not match:
        return None
    try:
        return date(*(int(part) for part in match.groups())).isoformat()
    except ValueError:
        return None


def read_case_rows(path: Path) -> list[CaseRow]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.worksheets[0]
    all_rows = list(sheet.iter_rows(values_only=True))
    if len(all_rows) < 2:
        raise RuntimeError("case workbook has no header row")
    headers = [clean(value) for value in all_rows[1]]
    required = [
        "文章名", "命名", "案例来源（链接）", "关联地区-省", "关联地区-市", "关联地区-区",
        "主体名", "是否为大学生创新创业", "大类", "小类", "发布日期", "100-300字摘要",
        "用到的AI工具或能力", "结果（效果、成果、数据）",
    ]
    missing = [name for name in required if name not in headers]
    if missing:
        raise RuntimeError(f"case workbook missing headers: {missing}")
    # This workbook retains a blank "关联地区-国家" header, while the data
    # rows use the original 14-column layout and place the province in column 4.
    # Prefer the observed row layout over the misleading header offset.
    first_data = all_rows[2] if len(all_rows) > 2 else ()
    legacy_layout = (
        len(first_data) >= 14
        and headers[3] == "关联地区-国家"
        and clean(first_data[3]) in {"北京市", "天津市", "上海市", "重庆市", "河北省", "山西省", "辽宁省", "吉林省", "黑龙江省", "江苏省", "浙江省", "安徽省", "福建省", "江西省", "山东省", "河南省", "湖北省", "湖南省", "广东省", "海南省", "四川省", "贵州省", "云南省", "陕西省", "甘肃省", "青海省", "台湾省", "内蒙古自治区", "广西壮族自治区", "西藏自治区", "宁夏回族自治区", "新疆维吾尔自治区", "香港特别行政区", "澳门特别行政区"}
    )
    compact_layout = legacy_layout and not parse_date(first_data[10])
    if compact_layout:
        # Current supplemental workbook also omits the country and publish-date
        # values in data rows, leaving summary/tools/outcome at columns 11-13.
        positions = {
            "文章名": 0, "命名": 1, "案例来源（链接）": 2,
            "关联地区-省": 3, "关联地区-市": 4, "关联地区-区": 5,
            "主体名": 6, "是否为大学生创新创业": 7, "大类": 8, "小类": 9,
            "发布日期": -1, "100-300字摘要": 10, "用到的AI工具或能力": 11,
            "结果（效果、成果、数据）": 12,
        }
    elif legacy_layout:
        positions = {
            "文章名": 0, "命名": 1, "案例来源（链接）": 2,
            "关联地区-省": 3, "关联地区-市": 4, "关联地区-区": 5,
            "主体名": 6, "是否为大学生创新创业": 7, "大类": 8, "小类": 9,
            "发布日期": 10, "100-300字摘要": 11, "用到的AI工具或能力": 12,
            "结果（效果、成果、数据）": 13,
        }
    else:
        positions = {name: headers.index(name) for name in required}
    records: list[CaseRow] = []
    for excel_row, values in enumerate(all_rows[2:], start=3):
        get = lambda name: clean(values[positions[name]]) if positions[name] < len(values) else None
        title = get("命名")
        url = get("案例来源（链接）")
        if not title and not url:
            continue
        if not title or not url:
            raise RuntimeError(f"case row {excel_row} has empty title or source URL")
        record = CaseRow(
            excel_row=excel_row,
            article_title=get("文章名") or title,
            title=title,
            source_url=url,
            province=normalize_province(get("关联地区-省")),
            city=get("关联地区-市"),
            district=get("关联地区-区"),
            actor_name=get("主体名"),
            student_startup=get("是否为大学生创新创业"),
            category=get("大类") or "其他",
            subcategory=get("小类") or "其他",
            publish_date=parse_date(values[positions["发布日期"]]),
            summary=get("100-300字摘要") or "",
            ai_tools=get("用到的AI工具或能力"),
            outcome=get("结果（效果、成果、数据）"),
        )
        if not record.summary:
            raise RuntimeError(f"case row {excel_row} has empty summary")
        records.append(record)
    return records


def dedupe_case_rows(rows: list[CaseRow]) -> tuple[list[CaseRow], list[int]]:
    seen: set[tuple[str, str, str]] = set()
    unique: list[CaseRow] = []
    duplicates: list[int] = []
    for row in rows:
        key = (row.source_url, row.actor_name or "", row.title)
        if key in seen:
            duplicates.append(row.excel_row)
            continue
        seen.add(key)
        unique.append(row)
    return unique, duplicates


def read_university_rows(path: Path) -> list[UniversityRow]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    result: list[UniversityRow] = []
    configs = [
        ("03高校 OPC 竞赛与活动", "activity", "高校 OPC 竞赛与活动", "activity_id", "activity_name"),
        ("04高校 OPC 案例", "case", "高校 OPC 案例", "case_id", "case_name"),
    ]
    for sheet_name, record_type, type_label, id_key, name_key in configs:
        if sheet_name not in workbook.sheetnames:
            raise RuntimeError(f"university workbook missing sheet: {sheet_name}")
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [clean(value) for value in rows[0]]
        for row in rows[1:]:
            raw = {headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers)) if headers[i]}
            code = clean(raw.get(id_key))
            name = clean(raw.get(name_key))
            if not code and not name:
                continue
            if not code or not name:
                raise RuntimeError(f"{sheet_name} contains a row without code/name")
            if record_type == "activity":
                summary = clean(raw.get("result_summary"))
                date_original = clean(raw.get("activity_date"))
                institution = clean(raw.get("host_institution"))
                notes = clean(raw.get("notes"))
            else:
                summary = clean(raw.get("summary_100_300"))
                date_original = clean(raw.get("publish_date"))
                institution = clean(raw.get("institution_name"))
                notes = clean(raw.get("notes"))
            result.append(UniversityRow(
                record_type=record_type,
                record_type_label=type_label,
                record_code=code,
                record_name=name,
                institution_name=institution,
                province=clean(raw.get("province")),
                city=clean(raw.get("city")),
                district=clean(raw.get("district")),
                verification_status=clean(raw.get("verification_status")) or "待人工核验",
                evidence_grade=clean(raw.get("evidence_grade")),
                date_original=date_original,
                source_title=clean(raw.get("source_title")),
                source_url=clean(raw.get("source_url")),
                summary=summary,
                notes=notes,
                raw={key: clean(value) for key, value in raw.items()},
            ))
    return result


def sql(value: Any) -> str:
    if value is None or value == "":
        return "NULL"
    return "'" + str(value).replace("\\", "\\\\").replace("'", "''") + "'"


def region_lookup(row: CaseRow) -> str:
    province, city, district = sql(row.province), sql(row.city), sql(row.district)
    predicates = []
    if row.district and row.city and row.province:
        predicates.append(
            f"(r.level='district' AND r.name={district} AND p.name={city} AND gp.name={province})"
        )
    if row.city and row.province:
        predicates.append(f"(r.level='city' AND r.name={city} AND p.name={province})")
    if row.province:
        predicates.append(f"(r.level='province' AND r.name={province})")
    if not predicates:
        return "(SELECT id FROM regions WHERE level='country' AND name='中国' ORDER BY id LIMIT 1)"
    return (
        "(SELECT r.id FROM regions r "
        "LEFT JOIN regions p ON p.id=r.parent_id "
        "LEFT JOIN regions gp ON gp.id=p.parent_id "
        "WHERE " + " OR ".join(predicates) +
        " ORDER BY CASE r.level WHEN 'district' THEN 1 WHEN 'city' THEN 2 ELSE 3 END, r.id LIMIT 1)"
    )


def case_tags(row: CaseRow) -> list[str]:
    tags = [row.category, row.subcategory]
    if row.student_startup == "是":
        tags.append("大学生创新创业")
    result: list[str] = []
    for tag in tags:
        if tag and tag not in result:
            result.append(tag)
    return result


def generate_update(cases: list[CaseRow], universities: list[UniversityRow], case_name: str, university_name: str) -> str:
    lines = [
        "SET NAMES utf8mb4;",
        "USE opc_platform;",
        "START TRANSACTION;",
        "",
        "CREATE TABLE IF NOT EXISTS university_opc_records (",
        "  id BIGINT PRIMARY KEY AUTO_INCREMENT,",
        "  record_type VARCHAR(30) NOT NULL,",
        "  record_type_label VARCHAR(100) NOT NULL,",
        "  record_code VARCHAR(50) NOT NULL,",
        "  record_name VARCHAR(500) NOT NULL,",
        "  institution_name VARCHAR(1000) NULL,",
        "  province VARCHAR(100) NULL, city VARCHAR(100) NULL, district VARCHAR(200) NULL,",
        "  verification_status VARCHAR(30) NOT NULL DEFAULT 'pending',",
        "  evidence_grade VARCHAR(10) NULL, date_original VARCHAR(100) NULL,",
        "  source_title VARCHAR(1000) NULL, source_url TEXT NULL, summary_text TEXT NULL,",
        "  notes TEXT NULL, raw_details LONGTEXT NULL,",
        "  created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,",
        "  updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,",
        "  UNIQUE KEY uk_university_opc_type_code (record_type, record_code),",
        "  INDEX idx_university_opc_type (record_type), INDEX idx_university_opc_province (province),",
        "  INDEX idx_university_opc_status (verification_status)",
        ") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;",
        "",
    ]
    for index, row in enumerate(cases, 1):
        batch_note = f"import_batch={BATCH};excel_file={case_name};excel_row={row.excel_row}"
        lines += [
            f"-- CASE {index}: Excel row {row.excel_row}",
            "INSERT INTO sources (title, source_type, publisher, url, local_file, accessed_at, notes, status)",
            "SELECT " + ", ".join([sql(row.article_title), sql("news"), "NULL", sql(row.source_url), "NULL", sql(ACCESS_DATE), sql(batch_note), sql("published")]),
            f"WHERE NOT EXISTS (SELECT 1 FROM sources WHERE url={sql(row.source_url)});",
            f"SET @source_id := (SELECT id FROM sources WHERE url={sql(row.source_url)} ORDER BY id LIMIT 1);",
            f"SET @region_id := {region_lookup(row)};",
            "INSERT INTO case_items (title, article_title, region_id, category, subcategory, actor_name, source_id, summary, business_model, ai_tools, outcome, tags, original_url, local_file, accessed_at, status, reviewer)",
            "SELECT " + ", ".join([
                sql(row.title), sql(row.article_title), "@region_id", sql(row.category), sql(row.subcategory),
                sql(row.actor_name), "@source_id", sql(row.summary),
                sql("地区：" + " / ".join(part for part in [row.province, row.city, row.district] if part) + "\n案例类型：" + row.category + "\n小类：" + row.subcategory),
                sql(row.ai_tools), sql(row.outcome), sql(",".join(case_tags(row))), sql(row.source_url), "NULL", sql(ACCESS_DATE), sql("published"), sql(CASE_REVIEWER),
            ]),
            "WHERE @region_id IS NOT NULL AND @source_id IS NOT NULL",
            f"  AND NOT EXISTS (SELECT 1 FROM case_items WHERE original_url={sql(row.source_url)} AND COALESCE(actor_name,'')=COALESCE({sql(row.actor_name)},''));",
            f"SET @case_id := (SELECT id FROM case_items WHERE original_url={sql(row.source_url)} AND COALESCE(actor_name,'')=COALESCE({sql(row.actor_name)},'') ORDER BY id DESC LIMIT 1);",
        ]
        for tag in case_tags(row):
            industry = "1" if tag == row.category else "0"
            lines += [
                "INSERT INTO tags (name, tag_type, is_industry, sort_order)",
                f"VALUES ({sql(tag)}, 'case', {industry}, 0) ON DUPLICATE KEY UPDATE is_industry=GREATEST(is_industry, VALUES(is_industry)), id=LAST_INSERT_ID(id);",
                "SET @tag_id := LAST_INSERT_ID();",
                "INSERT IGNORE INTO case_tags (case_id, tag_id) VALUES (@case_id, @tag_id);",
            ]
        lines.append("")
    for row in universities:
        raw = json.dumps(row.raw, ensure_ascii=False, separators=(",", ":"))
        columns = [row.record_type, row.record_type_label, row.record_code, row.record_name, row.institution_name, row.province, row.city, row.district, row.verification_status, row.evidence_grade, row.date_original, row.source_title, row.source_url, row.summary, row.notes, raw]
        values = ",".join(sql(value) for value in columns)
        lines += [
            f"-- UNIVERSITY {row.record_type}:{row.record_code}",
            "INSERT INTO university_opc_records (record_type,record_type_label,record_code,record_name,institution_name,province,city,district,verification_status,evidence_grade,date_original,source_title,source_url,summary_text,notes,raw_details)",
            f"VALUES ({values})",
            "ON DUPLICATE KEY UPDATE record_type_label=VALUES(record_type_label),record_name=VALUES(record_name),institution_name=VALUES(institution_name),province=VALUES(province),city=VALUES(city),district=VALUES(district),verification_status=VALUES(verification_status),evidence_grade=VALUES(evidence_grade),date_original=VALUES(date_original),source_title=VALUES(source_title),source_url=VALUES(source_url),summary_text=VALUES(summary_text),notes=VALUES(notes),raw_details=VALUES(raw_details);",
        ]
    lines += ["", "COMMIT;", ""]
    return "\n".join(lines)


def generate_rollback() -> str:
    return """SET NAMES utf8mb4;
USE opc_platform;
START TRANSACTION;
DELETE rel FROM case_tags rel INNER JOIN case_items c ON c.id=rel.case_id WHERE c.reviewer='excel-import-20260929';
DELETE FROM case_items WHERE reviewer='excel-import-20260929';
DELETE FROM sources WHERE notes LIKE '%import_batch=data-import-20260929%' AND NOT EXISTS (SELECT 1 FROM case_items c WHERE c.source_id=sources.id) AND NOT EXISTS (SELECT 1 FROM policies p WHERE p.source_id=sources.id);
DELETE t FROM tags t LEFT JOIN case_tags ct ON ct.tag_id=t.id WHERE t.tag_type='case' AND t.name IN ('大学生创新创业') AND ct.id IS NULL;
DELETE FROM university_opc_records WHERE record_code LIKE 'TMP-A%' OR record_code LIKE 'TMP-K%';
COMMIT;
"""


def generate_postcheck() -> str:
    return """SET NAMES utf8mb4;
USE opc_platform;
SELECT CONCAT('case_batch=',COUNT(*)) FROM case_items WHERE reviewer='excel-import-20260929';
SELECT CONCAT('case_sources=',COUNT(DISTINCT source_id)) FROM case_items WHERE reviewer='excel-import-20260929';
SELECT CONCAT('case_missing_region=',COUNT(*)) FROM case_items WHERE reviewer='excel-import-20260929' AND region_id IS NULL;
SELECT CONCAT('case_missing_tags=',COUNT(*)) FROM case_items c LEFT JOIN case_tags ct ON ct.case_id=c.id WHERE c.reviewer='excel-import-20260929' AND ct.id IS NULL;
SELECT CONCAT('university_total=',COUNT(*)) FROM university_opc_records;
SELECT record_type,COUNT(*) FROM university_opc_records GROUP BY record_type ORDER BY record_type;
SELECT verification_status,COUNT(*) FROM university_opc_records GROUP BY verification_status ORDER BY verification_status;
SELECT CONCAT('university_missing_required=',COUNT(*)) FROM university_opc_records WHERE record_code IS NULL OR record_code='' OR record_name IS NULL OR record_name='' OR source_url IS NULL OR source_url='';
"""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--case-excel", type=Path, default=DEFAULT_CASE_XLSX)
    parser.add_argument("--university-excel", type=Path, default=DEFAULT_UNIVERSITY_XLSX)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    raw_cases = read_case_rows(args.case_excel)
    cases, duplicate_case_rows = dedupe_case_rows(raw_cases)
    universities = read_university_rows(args.university_excel)
    if len(raw_cases) != 35:
        raise RuntimeError(f"expected 35 case rows, got {len(raw_cases)}")
    if len(universities) != 14:
        raise RuntimeError(f"expected 14 university rows, got {len(universities)}")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    plan = {
        "batch": BATCH,
        "inputs": {
            "case_excel": str(args.case_excel),
            "case_sha256": hashlib.sha256(args.case_excel.read_bytes()).hexdigest(),
            "university_excel": str(args.university_excel),
            "university_sha256": hashlib.sha256(args.university_excel.read_bytes()).hexdigest(),
        },
        "counts": {"case_rows_input": len(raw_cases), "case_rows_unique": len(cases), "case_rows_deduplicated": len(duplicate_case_rows), "university_rows": len(universities), "university_activity": sum(r.record_type == "activity" for r in universities), "university_case": sum(r.record_type == "case" for r in universities)},
        "duplicate_case_excel_rows": duplicate_case_rows,
        "empty_university_sheets": ["01高校 OPC 社区", "02高校 OPC 支持政策"],
        "case_rows": [asdict(row) for row in cases],
        "university_rows": [asdict(row) for row in universities],
    }
    (args.output_dir / "import-plan.json").write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.output_dir / "update.sql").write_text(generate_update(cases, universities, args.case_excel.name, args.university_excel.name), encoding="utf-8")
    (args.output_dir / "rollback.sql").write_text(generate_rollback(), encoding="utf-8")
    (args.output_dir / "postcheck.sql").write_text(generate_postcheck(), encoding="utf-8")
    print(json.dumps(plan["counts"], ensure_ascii=False))
    print(args.output_dir.resolve())


if __name__ == "__main__":
    main()
