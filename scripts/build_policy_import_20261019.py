"""Build a guarded, idempotent SQL import from the 2026-10-19 policy workbook."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils.datetime import from_excel


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = Path(r"D:\qq文件\人工智能OPC政策汇总_1019新.xlsx")
OUTPUT_DIR = ROOT / "outputs" / "policy-import-20261019"
SHEET = "Sheet1"
REVIEWER = "manual-workbook-20261019"
BATCH = "policy-import-20261019"
ACCESS_DATE = "2026-10-22"
TOPIC_FIELDS = [
    ("具体政策·算力技术", "算力技术"),
    ("具体政策·财政激励", "财政激励"),
    ("具体政策·场景开放", "场景开放"),
    ("具体政策·创业生态", "创业生态"),
    ("具体政策·金融资本", "金融资本"),
    ("具体政策·制度治理", "制度治理"),
    ("具体政策·人才培育", "人才培育"),
]
MATERIALS = {
    "正式文件": ("formal_policy", "published"),
    "征求意见稿": ("consultation_draft", "consultation"),
    "标准规范": ("standard_reference", "published"),
    "官方平台服务信息": ("official_platform_service", "published"),
    "其他": ("nonformal_policy_info", "published"),
}
LEVELS = {
    "国家级": "national",
    "省级": "provincial",
    "市级": "city",
    "市级部门": "city",
    "区级": "district",
    "县级": "district",
    "功能区级": "district",
}
SCOPES = {"综合类": "comprehensive", "专项类": "other", "综合": "comprehensive", "专项": "other"}
EXPECTED_IDS = [158, 159, 160, 161, 162]


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def quote(value) -> str:
    if value in (None, ""):
        return "NULL"
    return "'" + str(value).replace("\\", "\\\\").replace("'", "''") + "'"


def sql_date(value) -> str:
    if isinstance(value, (int, float)):
        converted = from_excel(value)
        return quote(converted.strftime("%Y-%m-%d"))
    text = clean(value)
    match = re.search(r"\b(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})日?\b", text)
    if not match:
        return "NULL"
    return quote(f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}")


def load_records(path: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in workbook.sheetnames:
        raise RuntimeError(f"missing worksheet: {SHEET}")
    rows = workbook[SHEET].iter_rows(values_only=True)
    headers = [clean(value) for value in next(rows)]
    records = []
    for excel_row, values in enumerate(rows, start=2):
        row = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
        if row.get("政策id") is None:
            continue
        record_id = int(row["政策id"])
        nature_label = clean(row.get("资料性质"))
        scope_label = clean(row.get("是否综合类"))
        level_label = clean(row.get("政策等级"))
        topics_text = clean(row.get("涉及主题"))
        topics = [item.strip() for item in re.split(r"[、,，;；\s]+", topics_text) if item.strip()]
        measures = {label: clean(row.get(field)) for field, label in TOPIC_FIELDS if clean(row.get(field)) not in ("无", "无。", "—", "-")}
        if set(topics) != set(measures):
            raise RuntimeError(f"topic/measure mismatch at Excel row {excel_row}: {topics} != {list(measures)}")
        if nature_label not in MATERIALS or scope_label not in SCOPES or level_label not in LEVELS:
            raise RuntimeError(f"unsupported mapping at Excel row {excel_row}: {nature_label}/{scope_label}/{level_label}")
        status_text = clean(row.get("政策状态"))
        status = "expired" if status_text in {"失效", "已失效", "废止"} else MATERIALS[nature_label][1]
        item = {
            "workbookId": record_id,
            "excelRow": excel_row,
            "title": clean(row.get("政策标题")),
            "url": clean(row.get("政策原文网页链接")),
            "evidenceUrl": clean(row.get("辅证链接")) or None,
            "statusText": status_text,
            "status": status,
            "level": LEVELS[level_label],
            "province": clean(row.get("省")),
            "city": clean(row.get("市")) or None,
            "district": clean(row.get("区")) or None,
            "documentNo": clean(row.get("文号")) or None,
            "issuer": clean(row.get("发文单位")),
            "publishDate": sql_date(row.get("发布日期")),
            "effectiveDate": sql_date(row.get("开始实施时间")),
            "validPeriod": clean(row.get("政策有效时长")) or None,
            "repealDate": sql_date(row.get("废止日期")),
            "natureLabel": nature_label,
            "nature": MATERIALS[nature_label][0],
            "scopeLabel": scope_label,
            "policyType": SCOPES[scope_label],
            "summary": clean(row.get("summary摘要（100字左右）")),
            "topics": topics,
            "measures": measures,
            "other": clean(row.get("具体政策·其他")),
        }
        missing = [key for key in ("title", "url", "province", "issuer", "summary") if not item[key]]
        if missing:
            raise RuntimeError(f"missing required fields at Excel row {excel_row}: {missing}")
        records.append(item)
    records.sort(key=lambda item: item["workbookId"])
    if [item["workbookId"] for item in records] != EXPECTED_IDS:
        raise RuntimeError(f"expected workbook IDs {EXPECTED_IDS}, got {[item['workbookId'] for item in records]}")
    urls = [item["url"] for item in records]
    if len(urls) != len(set(urls)):
        raise RuntimeError("duplicate source URLs within workbook")
    return records


def make_key_points(item: dict) -> str:
    lines = [
        f"资料性质：{item['natureLabel']}",
        "OPC相关性：明确OPC",
        f"涉及主题：{'、'.join(item['topics'])}",
        f"政策范围：{item['scopeLabel']}",
        f"适用地区：{' / '.join(filter(None, [item['province'], item['city'], item['district']]))}",
        f"政策状态：{item['statusText']}",
    ]
    if item["repealDate"]:
        lines.append(f"废止日期：{item['repealDate'].strip(chr(39))}")
    return "\n".join(lines)


def generate(records: list[dict], workbook_name: str) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "import-review.json").write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = ["SET NAMES utf8mb4;", "START TRANSACTION;", f"-- {BATCH}; workbook IDs 158-162; Sheet2 is reference-only and is not imported."]
    for item in records:
        notes = f"import_batch={BATCH}; workbook_id={item['workbookId']}; workbook={workbook_name}; user-reviewed source workbook"
        support_parts = [f"【{topic}】\n{text}" for topic, text in item["measures"].items()]
        if item["other"] and item["other"] not in ("无", "无。", "—", "-"):
            support_parts.append(f"【其他说明】\n{item['other']}")
        support = "\n\n".join(support_parts)
        lines.extend([
            "",
            f"-- Workbook ID {item['workbookId']}: {item['title']} (Excel row {item['excelRow']})",
            "INSERT INTO sources (title,source_type,publisher,url,local_file,accessed_at,notes,status,ai_evidence_status,evidence_revision)",
            f"SELECT {quote(item['title'])},'government_site',{quote(item['issuer'])},{quote(item['url'])},NULL,{quote(ACCESS_DATE)},{quote(notes)},'published','verified',1",
            f"WHERE NOT EXISTS (SELECT 1 FROM sources WHERE url={quote(item['url'])});",
            f"SET @source_id := (SELECT id FROM sources WHERE url={quote(item['url'])} ORDER BY id LIMIT 1);",
            f"SET @region_id := (SELECT id FROM regions WHERE name={quote(item['province'])} AND level='province' ORDER BY id LIMIT 1);",
            "INSERT INTO policies (title,region_id,issuing_body,document_no,publish_date,effective_date,valid_period,source_id,policy_level,policy_type,material_nature,applicability_mode,summary,key_points,support_measures,tags,original_url,evidence_url,local_file,accessed_at,status,reviewer,ai_evidence_status,evidence_revision)",
            f"SELECT {quote(item['title'])},@region_id,{quote(item['issuer'])},{quote(item['documentNo'])},{item['publishDate']},{item['effectiveDate']},{quote(item['validPeriod'])},@source_id,{quote(item['level'])},{quote(item['policyType'])},{quote(item['nature'])},'unclassified',{quote(item['summary'])},{quote(make_key_points(item))},{quote(support)},{quote('，'.join(item['topics']))},{quote(item['url'])},{quote(item['evidenceUrl'])},NULL,{quote(ACCESS_DATE)},{quote(item['status'])},{quote(REVIEWER)},'verified',1",
            f"WHERE @source_id IS NOT NULL AND @region_id IS NOT NULL AND NOT EXISTS (SELECT 1 FROM policies WHERE original_url={quote(item['url'])});",
            f"SET @policy_id := (SELECT id FROM policies WHERE original_url={quote(item['url'])} ORDER BY id LIMIT 1);",
        ])
        for order, (_, topic) in enumerate(TOPIC_FIELDS, 1):
            if topic in item["topics"]:
                lines.extend([
                    f"INSERT INTO tags (name,tag_type,is_industry,sort_order) VALUES ({quote(topic)},'policy',0,{order}) ON DUPLICATE KEY UPDATE id=LAST_INSERT_ID(id),sort_order=VALUES(sort_order),is_industry=0;",
                    "SET @tag_id := LAST_INSERT_ID();",
                    "INSERT IGNORE INTO policy_tags (policy_id,tag_id) SELECT @policy_id,@tag_id WHERE @policy_id IS NOT NULL;",
                ])
    url_list = ",".join(quote(item["url"]) for item in records)
    lines += [
        "", "COMMIT;", "",
        f"SELECT CONCAT('batch_policies=',COUNT(*)) FROM policies WHERE reviewer={quote(REVIEWER)};",
        f"SELECT CONCAT('batch_sources=',COUNT(*)) FROM sources WHERE notes LIKE {quote('%import_batch=' + BATCH + '%')};",
        f"SELECT CONCAT('missing=',COUNT(*)) FROM policies WHERE reviewer={quote(REVIEWER)} AND (title='' OR issuing_body='' OR summary='' OR original_url='' OR region_id IS NULL OR source_id IS NULL OR tags='');",
        f"SELECT CONCAT('tagless=',COUNT(*)) FROM policies p WHERE p.reviewer={quote(REVIEWER)} AND NOT EXISTS (SELECT 1 FROM policy_tags pt WHERE pt.policy_id=p.id);",
        f"SELECT CONCAT('topic_relations=',COUNT(*),',topic_names=',COUNT(DISTINCT t.name)) FROM policy_tags pt JOIN policies p ON p.id=pt.policy_id JOIN tags t ON t.id=pt.tag_id WHERE p.reviewer={quote(REVIEWER)};",
        f"SELECT CONCAT('duplicate_urls=',COUNT(*)) FROM (SELECT original_url FROM policies WHERE original_url IN ({url_list}) GROUP BY original_url HAVING COUNT(*)>1) duplicates;",
        f"SELECT id,title,material_nature,status,original_url FROM policies WHERE reviewer={quote(REVIEWER)} ORDER BY id;",
    ]
    (OUTPUT_DIR / "update.sql").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps({
        "records": len(records),
        "materials": dict(Counter(item["natureLabel"] for item in records)),
        "statuses": dict(Counter(item["statusText"] for item in records)),
        "topicRelations": sum(len(item["topics"]) for item in records),
        "workbookIds": [item["workbookId"] for item in records],
        "output": str(OUTPUT_DIR),
    }, ensure_ascii=False, indent=2))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    args = parser.parse_args()
    generate(load_records(args.source), args.source.name)


if __name__ == "__main__":
    main()
